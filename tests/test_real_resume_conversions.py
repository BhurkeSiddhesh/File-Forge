"""Conversion tests against a real, densely-formatted PDF (tests/fixtures/sample_resume.pdf).

WHY: every other converter test in this suite runs against PDFs generated at test
time via reportlab.pdfgen.canvas — single drawString calls with no bold/italic runs,
no bullet lists, no right-aligned columns, no hyperlinks. That's great for fast,
hermetic, isolated-scenario coverage (password handling, OCR fallback, page counts),
but it can't catch regressions that only show up on a real document: text reordering
across visually-aligned columns, or a converter simply choking on real-world
complexity. This file is additive to, not a replacement for, the synthetic-fixture
tests elsewhere in this suite.

The fixture is a real 2-page resume with bold section headers, an italicized summary,
bullet lists, and right-aligned date columns sitting on the same line as bold job
titles — exactly the kind of layout that can scramble text-extraction order in ways
single-drawString PDFs never exercise.
"""
from pathlib import Path

from docx import Document

from scripts.pdf_utils import (
    pdf_to_docx,
    pdf_to_excel,
    pdf_to_pptx,
    pdf_to_epub,
    extract_text_from_pdf,
    compress_pdf,
)

# Distinctive section headers from the real document (content markers, not PII).
SECTION_HEADERS = ["SUMMARY", "TECHNICAL SKILLS", "PROFESSIONAL EXPERIENCE", "EDUCATION"]


def test_pdf_to_docx(real_resume_pdf, tmp_path):
    output_path = pdf_to_docx(str(real_resume_pdf), str(tmp_path))
    assert Path(output_path).exists()
    assert output_path.endswith(".docx")

    doc = Document(output_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    for header in SECTION_HEADERS:
        assert header in text
    # Reading-order sanity: a real document's flush-left title / flush-right date
    # columns must not scramble extraction order.
    assert text.index("Fractal Analytics") < text.index("Think & Learn")


def test_pdf_to_excel_reports_no_tables(real_resume_pdf, tmp_path):
    """A resume has no real tables. This dense body-text layout used to
    trigger a false positive in the borderless-table fallback (PyMuPDF's
    "text" strategy invents column boundaries at recurring word-start
    x-coordinates), shredding paragraphs into garbled multi-column rows
    instead of correctly reporting zero tables — see the fill-consistency
    check in _extract_borderless_tables."""
    result = pdf_to_excel(str(real_resume_pdf), str(tmp_path))
    assert Path(result["output_path"]).exists()
    assert result["output_path"].endswith(".xlsx")
    assert result["tables_found"] == 0

    import openpyxl
    wb = openpyxl.load_workbook(result["output_path"])
    assert wb.sheetnames == ["Text Content"]


def test_pdf_to_pptx(real_resume_pdf, tmp_path):
    from pptx import Presentation

    result = pdf_to_pptx(str(real_resume_pdf), str(tmp_path))
    assert Path(result).exists()
    assert result.endswith(".pptx")

    prs = Presentation(result)
    assert len(prs.slides) == 2


def test_pdf_to_epub(real_resume_pdf, tmp_path):
    from ebooklib import epub

    result = pdf_to_epub(str(real_resume_pdf), str(tmp_path))
    assert Path(result).exists()
    assert result.endswith(".epub")

    book = epub.read_epub(result)
    chapters = [item for item in book.get_items_of_type(9) if item.file_name.startswith("page_")]
    assert len(chapters) == 2

    all_content = b"".join(item.get_content() for item in chapters)
    for header in SECTION_HEADERS:
        assert header.encode() in all_content
    assert all_content.index(b"Fractal Analytics") < all_content.index(b"Think &amp; Learn")


def test_extract_text_from_pdf(real_resume_pdf, tmp_path):
    result = extract_text_from_pdf(str(real_resume_pdf), str(tmp_path))
    assert Path(result["output_path"]).exists()
    assert result["page_count"] == 2

    text = Path(result["output_path"]).read_text(encoding="utf-8")
    for header in SECTION_HEADERS:
        assert header in text
    assert text.index("Fractal Analytics") < text.index("Think & Learn")


def test_compress_pdf(real_resume_pdf, tmp_path):
    result = compress_pdf(str(real_resume_pdf), str(tmp_path))
    output_path = Path(result["output_path"])
    assert output_path.exists()
    assert output_path.suffix == ".pdf"

    # Still opens and still has both pages after compression.
    import fitz
    doc = fitz.open(str(output_path))
    try:
        assert len(doc) == 2
    finally:
        doc.close()
